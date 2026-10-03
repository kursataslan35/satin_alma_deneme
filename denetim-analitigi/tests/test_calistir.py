from openpyxl import load_workbook

from calistir import main


def _csv(tmp_path, satirlar):
    dosya = tmp_path / "bordro.csv"
    baslik = "employee_id,department,position,salary_system,salary_received,salary_difference,fraud_flag"
    dosya.write_text(baslik + "\n" + "\n".join(satirlar) + "\n", encoding="utf-8")
    return dosya


def test_uctan_uca_rapor_uretir(tmp_path):
    dosya = _csv(tmp_path, ["E1,IT,Staff,100,100,0,0", "E2,HR,Staff,100,80,20,1", "E3,HR,Staff,100,90,10,0"])
    assert main(["bordro_maas_farki", str(dosya), "--cikti", str(tmp_path / "rapor")]) == 0
    rapor = next((tmp_path / "rapor").glob("*.xlsx"))
    kitap = load_workbook(rapor)
    assert kitap.sheetnames == ["Özet", "Profil", "Kalite", "İstisnalar", "Doğrulama"]
    ozet = {satir[0]: satir[1] for satir in kitap["Özet"].iter_rows(min_row=2, values_only=True)}
    assert ozet["İstisna sayısı"] == 2 and ozet["İsabet %"] == 50 and ozet["Yakalama %"] == 100


def test_kapi_kapaliysa_senaryo_calismaz_cikis_kodu_2(tmp_path):
    dosya = _csv(tmp_path, ["E1,IT,Staff,100,bilinmiyor,0,0"])
    assert main(["bordro_maas_farki", str(dosya), "--cikti", str(tmp_path / "rapor")]) == 2
    kitap = load_workbook(next((tmp_path / "rapor").glob("*.xlsx")))
    assert "İstisnalar" not in kitap.sheetnames
